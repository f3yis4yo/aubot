"""UC02 document validation, text extraction and routing.

Validation runs in a fixed order, cheapest check first, so invalid files are
rejected in milliseconds (TC02: error within 5 s):
  empty -> too large -> extension -> file signature (magic bytes) -> parse.

The file signature check means a renamed .exe is rejected even if it is called
"invoice.pdf". DOCX files are ZIP archives, so the total uncompressed size is
checked before parsing to defuse zip bombs.

Scanned PDFs with no text layer are accepted and flagged for manual review
(Assessment 1: "a known limit, not a Sprint 3 failure").
"""
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field

from django.conf import settings

MAX_PDF_PAGES_PARSED = 50
MAX_DOCX_UNCOMPRESSED = 60 * 1024 * 1024
EXCERPT_CHARS = 2000


class DocumentRejected(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass
class ParsedDocument:
    file_type: str
    sha256: str
    size_bytes: int
    text: str
    page_count: int | None
    has_text_layer: bool
    word_count: int = 0
    category: str = "general"
    department: str = "Customer Care"
    routing_reason: str = ""
    matched_keywords: list[str] = field(default_factory=list)


def _limit_mb() -> float:
    return settings.AUBOT["MAX_UPLOAD_BYTES"] / (1024 * 1024)


def check_size(size: int):
    if size <= 0:
        raise DocumentRejected("empty", "The file is empty. Please choose a PDF or Word document that has content.")
    if size > settings.AUBOT["MAX_UPLOAD_BYTES"]:
        raise DocumentRejected(
            "too_large", f"The file is {size / (1024 * 1024):.1f} MB. The limit is {_limit_mb():.0f} MB; please compress or split it."
        )


def _file_type(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in {"pdf", "docx"}:
        raise DocumentRejected("unsupported_type", "Only PDF (.pdf) and Word (.docx) files are accepted.")
    if ext == "pdf" and not data[:1024].lstrip().startswith(b"%PDF-"):
        raise DocumentRejected("corrupted", "This file is not a valid PDF. It may be damaged or renamed from another format.")
    if ext == "docx" and not data.startswith(b"PK\x03\x04"):
        raise DocumentRejected("corrupted", "This file is not a valid Word document. It may be damaged or renamed.")
    return ext


def _parse_pdf(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentRejected("encrypted", "This PDF is password-protected. Please upload an unlocked copy.")
        pages = reader.pages
        text = "\n".join((pages[i].extract_text() or "") for i in range(min(len(pages), MAX_PDF_PAGES_PARSED)))
        return text, len(pages)
    except DocumentRejected:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError, OSError, RecursionError) as exc:
        raise DocumentRejected("corrupted", "This PDF could not be read; it appears to be damaged.") from exc


def _parse_docx(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = zf.namelist()
            if "word/document.xml" not in names:
                raise DocumentRejected("corrupted", "This file is not a valid Word document.")
            if sum(info.file_size for info in zf.infolist()) > MAX_DOCX_UNCOMPRESSED:
                raise DocumentRejected("corrupted", "This Word document is malformed (it expands to an unsafe size).")
    except zipfile.BadZipFile as exc:
        raise DocumentRejected("corrupted", "This Word document could not be opened; it appears to be damaged.") from exc

    import docx

    try:
        doc = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises a variety of lxml/KeyError types
        raise DocumentRejected("corrupted", "This Word document could not be read; it appears to be damaged.") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(p for p in parts if p.strip())


# (category, department, keywords). Scored by keyword hits in filename + text.
ROUTES = [
    ("insurance_claim", "Claims & Insurance", ["insurance", "insurer", "claim", "policy number", "excess", "assessor", "storm damage"]),
    ("invoice_billing", "Accounts", ["invoice", "tax invoice", "amount due", "remittance", "payment", "receipt", "abn", "bsb"]),
    ("quote_acceptance", "Sales & Scheduling", ["quote", "quotation", "estimate", "accept", "signed", "signature", "booking"]),
    ("strata_approval", "Compliance", ["strata", "owners corporation", "body corporate", "by-law", "bylaw", "approval", "landlord", "consent"]),
    ("compliance_certificate", "Compliance", ["certificate", "compliance", "ccew", "safety check", "inspection report", "smoke alarm"]),
    ("plans_specs", "Estimating", ["floor plan", "plans", "drawing", "specification", "layout", "renovation", "architect"]),
    ("complaint", "Customer Care", ["complaint", "dissatisfied", "unacceptable", "damage caused", "refund", "ombudsman"]),
]


def route(filename: str, text: str) -> tuple[str, str, str, list[str]]:
    haystack = f"{filename} {text[:20000]}".lower()
    best = ("general", "Customer Care", [])
    best_score = 0
    for category, department, keywords in ROUTES:
        hits = [k for k in keywords if re.search(rf"\b{re.escape(k)}", haystack)]
        if len(hits) > best_score:
            best, best_score = (category, department, hits), len(hits)
    category, department, hits = best
    reason = f"Matched: {', '.join(hits)}" if hits else "No routing keywords found; sent to general enquiries"
    return category, department, reason, hits


def process(filename: str, data: bytes) -> ParsedDocument:
    """Validate, parse and route. Raises DocumentRejected with a user-facing message."""
    check_size(len(data))
    file_type = _file_type(filename, data)
    if file_type == "pdf":
        text, pages = _parse_pdf(data)
    else:
        text, pages = _parse_docx(data), None
    words = len(text.split())
    if file_type == "docx" and words == 0:
        raise DocumentRejected("empty", "This Word document has no text in it. Please check you chose the right file.")
    category, department, reason, hits = route(filename, text)
    return ParsedDocument(
        file_type=file_type,
        sha256=hashlib.sha256(data).hexdigest(),
        size_bytes=len(data),
        text=text,
        page_count=pages,
        has_text_layer=words > 0,
        word_count=words,
        category=category,
        department=department,
        routing_reason=reason if words or file_type == "docx" else "Scanned PDF without a text layer; needs manual review",
        matched_keywords=hits,
    )
