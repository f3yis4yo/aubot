"""Synthetic TC02 test files, generated in memory so the suite is reproducible
and no customer documents are needed."""
import io
import os
import random


def make_pdf(text: str | None, pad_bytes: int = 0) -> bytes:
    """A valid single-page PDF. text=None gives a page with no text layer (like a scan).
    pad_bytes adds an unreferenced binary stream to reach a target file size."""
    escaped = (text or "").replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = f"BT /F1 11 Tf 72 720 Td ({escaped}) Tj ET".encode() if text else b"0 0 m 100 100 l S"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    if pad_bytes:
        blob = os.urandom(pad_bytes)
        objects.append(b"<< /Length %d >>\nstream\n" % len(blob) + blob + b"\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def make_docx(paragraphs: list[str], approx_size: int = 0) -> bytes:
    import docx

    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if approx_size:
        rng = random.Random(42)
        # Random reference codes compress poorly, so the file reaches the target size.
        while True:
            for _ in range(400):
                document.add_paragraph(" ".join(f"{rng.getrandbits(64):016x}" for _ in range(12)))
            buf = io.BytesIO()
            document.save(buf)
            if buf.tell() >= approx_size:
                return buf.getvalue()
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


MB = 1024 * 1024


def tc02_cases() -> list[dict]:
    """Assessment 2 TC02: valid PDF (2 MB), valid DOCX (1 MB), oversized (8 MB), plus the
    exception cases listed in Section 7.3: unsupported, empty, corrupted, renamed, scanned."""
    claim = "Insurance claim form. Policy number HX-449201. Insurer: Coastal Mutual. Storm damage to roof plumbing."
    return [
        {"name": "valid_claim.pdf", "data": make_pdf(claim, pad_bytes=2 * MB), "expect": "accepted", "label": "Valid PDF (2 MB)"},
        {
            "name": "signed_quote.docx",
            "data": make_docx(["Quote acceptance", "I accept quotation Q-1182 and have signed below."], approx_size=1 * MB),
            "expect": "accepted",
            "label": "Valid DOCX (1 MB)",
        },
        {"name": "scanned_letter.pdf", "data": make_pdf(None), "expect": "accepted", "label": "Scanned PDF (no text layer)"},
        {"name": "huge_plans.pdf", "data": make_pdf("Floor plan", pad_bytes=8 * MB), "expect": "too_large", "label": "Oversized PDF (8 MB)"},
        {"name": "notes.txt", "data": b"Please call me back", "expect": "unsupported_type", "label": "Unsupported format (.txt)"},
        {"name": "empty.pdf", "data": b"", "expect": "empty", "label": "Empty upload"},
        {"name": "broken.docx", "data": b"PK\x03\x04" + os.urandom(4096), "expect": "corrupted", "label": "Corrupted DOCX"},
        {"name": "invoice.pdf", "data": b"MZ\x90\x00" + os.urandom(2048), "expect": "corrupted", "label": "Executable renamed to .pdf"},
        {"name": "blank.docx", "data": make_docx([]), "expect": "empty", "label": "DOCX with no text"},
    ]
